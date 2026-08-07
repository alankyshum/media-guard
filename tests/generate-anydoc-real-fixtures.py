#!/usr/bin/env python3
"""Generate small, structurally-valid office fixtures for the real anydoc E2E probe."""
from pathlib import Path
import sys
import zipfile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches, Pt
from odf import text, table, teletype
from odf.opendocument import OpenDocumentPresentation, OpenDocumentSpreadsheet, OpenDocumentText
from odf.style import Style, TextProperties

root = Path(sys.argv[1] if len(sys.argv) > 1 else "/tmp/anydoc-real")
root.mkdir(parents=True, exist_ok=True)

heading = "Quarterly Report"
bold = "MARKER-BOLD-7391"
table_values = (("Alpha", "Beta", "Gamma"), ("Delta", "Epsilon", "Zeta"))
bullets = ("MARKER-BULLET-1842", "MARKER-LIST-5930")

def add_table(doc):
    table_ = doc.add_table(rows=2, cols=3)
    for r, row in enumerate(table_values):
        for c, value in enumerate(row): table_.cell(r, c).text = value

doc = Document()
doc.add_heading(heading, level=1)
p = doc.add_paragraph(); p.add_run("Important finding: "); p.add_run(bold).bold = True
add_table(doc)
for item in bullets: doc.add_paragraph(item, style="List Bullet")
doc.save(root / "quarterly-report.docx")

wb = Workbook(); ws = wb.active; ws.title = "Quarterly"
ws["A1"] = heading; ws["A2"] = bold
for r, row in enumerate(table_values, 4):
    for c, value in enumerate(row, 1): ws.cell(r, c, value)
ws["A7"] = bullets[0]; ws["A8"] = bullets[1]
wb.save(root / "quarterly-report.xlsx")

prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[5])
box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(5))
tf = box.text_frame; p = tf.paragraphs[0]; p.text = heading; p.font.size = Pt(32)
p = tf.add_paragraph(); p.text = bold; p.font.bold = True
p = tf.add_paragraph(); p.text = " | ".join(table_values[0])
for item in bullets:
    p = tf.add_paragraph(); p.text = item; p.level = 0
prs.save(root / "quarterly-report.pptx")

def odf_paragraph(document, value, bold_text=False):
    p = text.P(text=value)
    document.text.addElement(p)
    return p
def odf_table(document):
    t = table.Table(name="MarkerTable")
    for row in table_values:
        tr = table.TableRow()
        for value in row:
            cell = table.TableCell(); cell.addElement(text.P(text=value)); tr.addElement(cell)
        t.addElement(tr)
    document.text.addElement(t)

odt = OpenDocumentText(); odf_paragraph(odt, heading); odf_paragraph(odt, bold); odf_table(odt)
for item in bullets: odf_paragraph(odt, item)
odt.save(str(root / "quarterly-report.odt"))

ods = OpenDocumentSpreadsheet(); sheet = table.Table(name="Quarterly")
for values in ((heading,), (bold,), *table_values, *[(x,) for x in bullets]):
    row = table.TableRow()
    for value in values:
        cell = table.TableCell(); cell.addElement(text.P(text=value)); row.addElement(cell)
    sheet.addElement(row)
ods.spreadsheet.addElement(sheet); ods.save(str(root / "quarterly-report.ods"))

odp = OpenDocumentPresentation(); page = __import__('odf.draw', fromlist=['Page']).Page(name="page1", masterpagename="Default")
frame = __import__('odf.draw', fromlist=['Frame']).Frame(width="20cm", height="10cm", x="1cm", y="1cm")
box = __import__('odf.draw', fromlist=['TextBox']).TextBox()
for item in (heading, bold, *table_values[0], *bullets): box.addElement(text.P(text=item))
frame.addElement(box); page.addElement(frame); odp.presentation.addElement(page); odp.save(str(root / "quarterly-report.odp"))

(root / "quarterly-report.rtf").write_text(r"{\rtf1\ansi\deff0 {\b Quarterly Report}\par Important finding: {\b MARKER-BOLD-7391}\par \trowd\cellx2000\cellx4000\cellx6000 Alpha\cell Beta\cell Gamma\cell\row\trowd\cellx2000\cellx4000\cellx6000 Delta\cell Epsilon\cell Zeta\cell\row\par \bullet MARKER-BULLET-1842\par \bullet MARKER-LIST-5930\par}", encoding="ascii")
(root / "quarterly-report.csv").write_text("Quarterly Report,MARKER-BOLD-7391\nAlpha,Beta,Gamma\nDelta,Epsilon,Zeta\nMARKER-BULLET-1842,MARKER-LIST-5930\n", encoding="utf-8")

epub = root / "quarterly-report.epub"
with zipfile.ZipFile(epub, "w") as z:
    z.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
    z.writestr("META-INF/container.xml", """<?xml version='1.0'?><container version='1.0' xmlns='urn:oasis:names:tc:opendocument:xmlns:container'><rootfiles><rootfile full-path='OEBPS/content.opf' media-type='application/oebps-package+xml'/></rootfiles></container>""")
    z.writestr("OEBPS/content.opf", """<?xml version='1.0' encoding='utf-8'?><package version='3.0' xmlns='http://www.idpf.org/2007/opf' unique-identifier='bookid'><metadata xmlns:dc='http://purl.org/dc/elements/1.1/'><dc:identifier id='bookid'>marker-book</dc:identifier><dc:title>Quarterly Report</dc:title><dc:language>en</dc:language></metadata><manifest><item id='chapter' href='chapter.xhtml' media-type='application/xhtml+xml'/></manifest><spine><itemref idref='chapter'/></spine></package>""")
    z.writestr("OEBPS/chapter.xhtml", """<?xml version='1.0' encoding='utf-8'?><html xmlns='http://www.w3.org/1999/xhtml'><body><h1>Quarterly Report</h1><p><strong>MARKER-BOLD-7391</strong></p><table><tr><td>Alpha</td><td>Beta</td><td>Gamma</td></tr><tr><td>Delta</td><td>Epsilon</td><td>Zeta</td></tr></table><ul><li>MARKER-BULLET-1842</li><li>MARKER-LIST-5930</li></ul></body></html>""")
print(root)
